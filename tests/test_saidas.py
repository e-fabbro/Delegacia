"""F5: `matrizes`, `render` completo (md/xlsx/html, ponteiros), `handoff` e scripts de operação."""
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

import docx
import openpyxl
import pytest

from agencia import banco, caso, grafo, handoff, ingestao, matrizes, render, rif
from tests.fixtures.gerar import PF1, PF2, PJ1, CONTAS

FIXTURES = Path(__file__).resolve().parent / "fixtures"
RAIZ = Path(__file__).resolve().parent.parent
CCS, RIF, SIMBA = "DOC-001", "DOC-002", "DOC-003"
RE_TOKEN = re.compile(r"\b(PF|PJ|CT|TEL|EML|PIX|END)-\d{4}\b")

PRODUTO = """# Informação de Análise — TESTE

## 1. Objeto
Movimentação de PF-0001 e PF-0002 [F:DOC-003:tx#2].

## 4. Análise por fonte
- **FATO** PF-0001 transferiu R$ 20.000,00 a PF-0002 em 07/01/2026 [F:DOC-003:tx#2].
- *INFERÊNCIA* A conta CT-0002 tem perfil de passagem [F:DOC-003:agg#banco_passagem_horas48].

| entidade | papel | fonte |
|---|---|---|
| PF-0001 | titular | [F:DOC-002:com#1] |
| PJ-0001 | remetente | [F:DOC-002:com#1] |

> Segundo o comunicante, trata-se de valores incompatíveis com a renda [F:DOC-002:com#1].

1. Diligência BAN — PF-0003.
"""

ACHADO = {"id": "BAN-001", "agente": "analista-bancario", "rotulo": "FATO", "relevancia": "alta",
          "enunciado": "PF-0001 transferiu R$ 20.000,00 a PF-0002 em 07/01/2026.", "entidades": ["PF-0001", "PF-0002"],
          "valores": [{"valor": 20000.0, "origem": "banco lancamentos"}], "fontes": [{"doc_id": SIMBA, "localizador": "tx#2"}]}
HIPOTESE = {"id": "RIF-003", "agente": "analista-rif", "rotulo": "HIPOTESE", "relevancia": "media",
            "enunciado": "PF-0003 pode operar como conta de passagem.", "entidades": ["PF-0003"],
            "fontes": [{"doc_id": RIF, "localizador": "com#8"}], "diligencia": {"tipo": "BAN", "alvo": "PF-0003", "motivo": "confirmar passagem"}}


def cli(*args):
    return subprocess.run([sys.executable, "-m", "agencia", *args], capture_output=True, text=True)


@pytest.fixture
def caso_completo(raiz_casos):
    caso.novo("TESTE")
    d = raiz_casos / "TESTE"
    for nome in ("ccs_3contas.xlsx", "rif_12_sintetico.pdf", "simba_3contas.csv"):
        shutil.copy(FIXTURES / nome, d / "00_brutos" / nome)
    ingestao.ingerir("TESTE")
    banco.importar("TESTE", CCS)
    banco.importar("TESTE", SIMBA)
    rif.parse("TESTE", RIF)
    assert cli("banco", "passagem", "TESTE", "--salvar").returncode == 0
    for agente, a in (("analista-bancario", ACHADO), ("analista-rif", HIPOTESE)):
        pasta = d / "03_analises" / agente
        pasta.mkdir(parents=True)
        (pasta / "achados.jsonl").write_text(json.dumps(a) + "\n", encoding="utf-8")
    grafo.construir("TESTE")
    grafo.exportar("TESTE")
    (d / "04_produtos" / "informacao_analise_v1.md").write_text(PRODUTO, encoding="utf-8")
    return d


# ---------- matrizes ----------

def test_matrizes(caso_completo):
    r = matrizes.matrizes("TESTE")
    assert r["arquivo"] == "04_produtos/matrizes.xlsx"
    for nome in ("Documentos", "Contas", "Lancamentos", "Integridade", "Contrapartes", "Especie", "Fracionamento", "Passagem", "Circularidade",
                 "Cruzamentos", "LinhaTempo_Banco", "RIF_Comunicacoes", "RIF_Envolvidos", "RIF_Sobreposicao", "RIF_Consolidado", "Vinculos",
                 "Centrais", "LinhaTempo_Integrada", "Marcos", "Achados", "Diligencias"):
        assert r["planilhas"].get(nome), nome
    assert r["omitidas"] == []
    assert r["planilhas"]["Lancamentos"] == 32 and r["planilhas"]["RIF_Comunicacoes"] == 12 and r["planilhas"]["Achados"] == 2
    wb = openpyxl.load_workbook(caso_completo / "04_produtos" / "matrizes.xlsx")
    ws = wb["Lancamentos"]
    assert ws["A1"].value == "ponteiro" and ws["A1"].font.bold and ws.freeze_panes == "A2"
    tudo = " ".join(str(c.value) for w in wb.worksheets for row in w.iter_rows() for c in row if c.value is not None)
    for claro in (PF1["nome"], PF1["cpf"], PF2["cpf"], PJ1["cnpj"], CONTAS["A"]["conta"]):
        assert claro not in tudo
    assert "PF-0001" in tudo


def test_matrizes_parcial(raiz_casos):
    caso.novo("SO_RIF")
    shutil.copy(FIXTURES / "rif_12_sintetico.pdf", raiz_casos / "SO_RIF" / "00_brutos" / "rif_12_sintetico.pdf")
    ingestao.ingerir("SO_RIF")
    rif.parse("SO_RIF", "DOC-001")
    r = matrizes.matrizes("SO_RIF")
    assert "RIF_Comunicacoes" in r["planilhas"] and "Lancamentos" in r["omitidas"] and "Contas" in r["omitidas"]
    caso.novo("VAZIO")
    with pytest.raises(caso.ErroCaso):
        matrizes.matrizes("VAZIO")


# ---------- render ----------

def _texto_docx(p: Path) -> str:
    doc = docx.Document(p)
    partes = [x.text for x in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells]
    partes += [x.text for s in doc.sections for x in list(s.header.paragraphs) + list(s.footer.paragraphs)]
    return "\n".join(partes)


def test_criterio_aceite_7_docx_sem_tokens(caso_completo):
    r = render.render("TESTE", "informacao_analise_v1.md")
    assert r["ok"] and r["tokens_remanescentes"] == 0 and r["tokens_substituidos"] >= 8
    tudo = _texto_docx(Path(r["docx"]))
    assert not RE_TOKEN.search(tudo)
    assert PF1["nome"] in tudo and PF1["cpf"] in tudo and PF2["nome"] in tudo and PJ1["cnpj"] in tudo
    assert CONTAS["B"]["conta"].replace("-", "") in tudo.replace("-", "")            # CT-0002 reidentificada
    assert "SIGILOSO" in tudo and "Caso TESTE" in tudo
    assert "[DOC-003, lançamento 2]" in tudo and "[DOC-002, comunicação 1]" in tudo and "análise banco_passagem_horas48" in tudo
    assert "[F:DOC" not in tudo and "**" not in tudo
    doc = docx.Document(r["docx"])
    estilos = {p.style.name for p in doc.paragraphs}
    assert {"Heading 1", "Heading 2", "List Bullet", "List Number", "Intense Quote"} <= estilos
    negritos = [run.text for p in doc.paragraphs for run in p.runs if run.bold]
    assert "FATO" in negritos
    assert len(doc.tables) == 1 and doc.tables[0].rows[0].cells[0].paragraphs[0].runs[0].bold


def test_render_ponteiros_manter_e_remover(caso_completo):
    r = render.render("TESTE", "informacao_analise_v1.md", ponteiros="manter", docx_=False)
    md = Path(r["md"]).read_text(encoding="utf-8")
    assert "[F:DOC-003:tx#2]" in md and "docx" not in r
    r = render.render("TESTE", "informacao_analise_v1.md", ponteiros="remover")
    md = Path(r["md"]).read_text(encoding="utf-8")
    assert "[F:" not in md and "[DOC-" not in md and "em 07/01/2026." in md
    with pytest.raises(caso.ErroCaso):
        render.render("TESTE", "informacao_analise_v1.md", ponteiros="x")


def test_render_xlsx_e_html(caso_completo):
    matrizes.matrizes("TESTE")
    r = render.render("TESTE", "matrizes.xlsx")
    assert r["ok"] and r["xlsx"].endswith("render/matrizes.xlsx")
    wb = openpyxl.load_workbook(r["xlsx"])
    tudo = " ".join(str(c.value) for w in wb.worksheets for row in w.iter_rows() for c in row if isinstance(c.value, str))
    assert not RE_TOKEN.search(tudo) and PF1["nome"] in tudo and "[DOC-003, lançamento" in tudo
    r = render.render("TESTE", "grafo.html")
    html = Path(r["html"]).read_text(encoding="utf-8")
    assert r["ok"] and not RE_TOKEN.search(html) and PF1["nome"] in html
    with pytest.raises(caso.ErroCaso):
        render.render("TESTE", "render/matrizes.xlsx")
    (caso_completo / "04_produtos" / "x.bin").write_bytes(b"\x00")
    with pytest.raises(caso.ErroCaso):
        render.render("TESTE", "x.bin")


# ---------- handoff ----------

def test_handoff(caso_completo):
    r = handoff.handoff("TESTE")
    assert r["arquivo"] == "04_produtos/handoff.json" and re.fullmatch(r"[0-9a-f]{64}", r["sha256"])
    pacote = json.loads((caso_completo / "04_produtos" / "handoff.json").read_text(encoding="utf-8"))
    assert handoff.validar(pacote) == []
    assert pacote["caso"]["codinome"] == "TESTE" and pacote["produto"]["arquivo"] == "04_produtos/informacao_analise_v1.md" and pacote["produto"]["versao"] == 1
    assert [d["doc_id"] for d in pacote["documentos"]] == [CCS, RIF, SIMBA]
    simba = pacote["documentos"][2]
    assert (simba["periodo_inicio"], simba["periodo_fim"]) == ("2026-01-05", "2026-06-30")
    assert pacote["documentos"][1]["periodo_inicio"] == "2026-01-01"
    env = {e["pseudonimo"]: e for e in pacote["envolvidos"]}
    assert env["PF-0001"]["contas"] == ["CT-0001"] and "titular_de_conta" in env["PF-0001"]["papeis"] and "titular_rif" in env["PF-0001"]["papeis"]
    assert set(env["PF-0001"]["fontes"]) >= {"SIMBA", "RIF", "CCS"} and env["PF-0001"]["n_achados"] == 1 and env["PF-0001"]["relevancia_maxima"] == "alta"
    assert pacote["envolvidos"][0]["relevancia_maxima"] == "alta"                # ordenação por relevância
    assert pacote["achados"][0]["id"] == "BAN-001" and pacote["achados"][0]["fontes"] == [f"[F:{SIMBA}:tx#2]"]
    assert pacote["diligencias"] == [{"tipo": "BAN", "alvo": "PF-0003", "motivos": ["confirmar passagem"], "relevancia": "media", "achados": ["analista-rif:RIF-003"]}]
    assert pacote["resumo"] == {"documentos": 3, "por_tipo": {"CCS": 1, "RIF": 1, "SIMBA": 1}, "envolvidos": 7, "achados": 2,
                                "por_rotulo": {"FATO": 1, "HIPOTESE": 1}, "diligencias": 1, "convergencias": len(pacote["convergencias"])}
    assert any(c["pseudonimo"] == "PF-0002" for c in pacote["convergencias"])
    texto = json.dumps(pacote, ensure_ascii=False)
    for claro in (PF1["nome"], PF1["cpf"], PJ1["cnpj"]):
        assert claro not in texto


def test_handoff_reidentificado(caso_completo):
    r = handoff.handoff("TESTE", reidentificar=True)
    reid = (caso_completo / "04_produtos" / "render" / "handoff.json").read_text(encoding="utf-8")
    assert r["reidentificado"]["tokens_remanescentes"] == 0
    assert not RE_TOKEN.search(reid) and PF1["nome"] in reid and PF1["cpf"] in reid
    assert RE_TOKEN.search((caso_completo / "04_produtos" / "handoff.json").read_text(encoding="utf-8"))   # o pseudonimizado fica


def test_handoff_caso_vazio(raiz_casos):
    caso.novo("VAZIO")
    r = handoff.handoff("VAZIO")
    pacote = json.loads((raiz_casos / "VAZIO" / "04_produtos" / "handoff.json").read_text(encoding="utf-8"))
    assert pacote["produto"] is None and pacote["documentos"] == [] and pacote["resumo"]["achados"] == 0 and r["resumo"]["envolvidos"] == 0


def test_schema_handoff_rejeita_invalido(caso_completo):
    assert handoff.validar({"versao_handoff": "x"})           # faltam obrigatórios
    pacote = handoff.montar("TESTE")
    pacote["achados"][0]["fontes"] = ["DOC-003 tx#2"]         # ponteiro fora do formato
    pacote["envolvidos"][0]["pseudonimo"] = "MARIA"
    erros = handoff.validar(pacote)
    assert len(erros) == 2 and any("fontes" in e for e in erros) and any("pseudonimo" in e for e in erros)


# ---------- CLI e scripts ----------

def test_cli_saidas(caso_completo):
    r = cli("matrizes", "TESTE")
    assert r.returncode == 0 and json.loads(r.stdout)["planilhas"]["Achados"] == 2
    r = cli("render", "TESTE", "matrizes.xlsx", "--ponteiros", "manter")
    assert r.returncode == 0 and json.loads(r.stdout)["ok"] is True
    r = cli("render", "TESTE", "informacao_analise_v1.md", "--sem-docx")
    assert r.returncode == 0 and "docx" not in json.loads(r.stdout)
    r = cli("handoff", "TESTE", "--reidentificar", "--md")
    assert r.returncode == 0 and "handoff.json" in r.stdout


def test_scripts_de_operacao_sao_bash_valido():
    for rel in ("ops/teste_headless.sh", "hermes/nexo_run.sh", "ops/setup_vps.sh"):
        p = subprocess.run(["bash", "-n", str(RAIZ / rel)], capture_output=True, text=True)
        assert p.returncode == 0, (rel, p.stderr)
    texto = (RAIZ / "hermes" / "nexo_run.sh").read_text(encoding="utf-8")
    assert "sudo" in texto and "--output-format json" in texto and "nohup" in texto
