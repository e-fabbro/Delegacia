"""F1: `render` básico — reidentifica o produto fora do modelo e gera .docx/.md em 04_produtos/render/."""
import json
import re
import subprocess
import sys

import docx
import pytest

from agencia import caso, cofre, render
from tests.fixtures.gerar import PF1, PJ1

PRODUTO = """# Informação de Análise — TESTE

## Objeto
Movimentação de PF-0001 com PJ-0001 [F:DOC-001:p1].

| entidade | papel | valor |
|---|---|---|
| PF-0001 | titular | R$ 150.000,00 |
| PJ-0001 | remetente | R$ 150.000,00 |

- Contato: TEL-0001
- Pendente: CT-0001 sem extrato
"""


@pytest.fixture
def caso_com_produto(raiz_casos):
    caso.novo("TESTE")
    d = raiz_casos / "TESTE"
    cf = cofre.Cofre(d / "_cofre" / "identidades.db")
    cf.vincular("nome_pf", PF1["nome"], cf.pseudonimo("cpf", PF1["cpf"]))
    cf.vincular("nome_pj", PJ1["nome"], cf.pseudonimo("cnpj", PJ1["cnpj"]))
    cf.pseudonimo("telefone", PF1["tel"])
    cf.pseudonimo("conta", "1234/56789-0")
    (d / "04_produtos" / "informacao_analise_v1.md").write_text(PRODUTO, encoding="utf-8")
    return d


def test_render_gera_docx_sem_tokens(caso_com_produto):
    r = render.render("TESTE", "informacao_analise_v1.md")
    saida = caso_com_produto / "04_produtos" / "render" / "informacao_analise_v1.docx"
    assert saida.is_file() and r["docx"] == str(saida)
    doc = docx.Document(saida)
    textos = [p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells]
    tudo = "\n".join(textos)
    assert not re.search(r"\b(PF|PJ|CT|TEL|EML|PIX|END)-\d{4}\b", tudo)
    assert PF1["nome"] in tudo and PF1["cpf"] in tudo and PJ1["cnpj"] in tudo and PF1["tel"] in tudo
    assert "R$ 150.000,00" in tudo and "[F:DOC-001:p1]" in tudo
    assert doc.paragraphs[0].style.name.startswith("Heading")
    assert len(doc.tables) == 1 and len(doc.tables[0].rows) == 3


def test_render_gera_md_reidentificado(caso_com_produto):
    r = render.render("TESTE", "informacao_analise_v1.md")
    md = (caso_com_produto / "04_produtos" / "render" / "informacao_analise_v1.md").read_text(encoding="utf-8")
    assert r["md"].endswith("render/informacao_analise_v1.md")
    assert "PF-0001" not in md and PF1["nome"] in md
    assert r["tokens_substituidos"] >= 5 and r["tokens_remanescentes"] == 0


def test_render_token_desconhecido_fica_e_conta(caso_com_produto):
    (caso_com_produto / "04_produtos" / "x.md").write_text("PF-0099 sem cadastro", encoding="utf-8")
    r = render.render("TESTE", "x.md")
    assert r["tokens_remanescentes"] == 1


def test_render_recusa_fora_de_04_produtos(caso_com_produto):
    with pytest.raises(caso.ErroCaso):
        render.render("TESTE", "../estado.json")
    with pytest.raises(caso.ErroCaso):
        render.render("TESTE", "nao_existe.md")


def test_cli_render(caso_com_produto):
    p = subprocess.run([sys.executable, "-m", "agencia", "render", "TESTE", "informacao_analise_v1.md"], capture_output=True, text=True)
    assert p.returncode == 0, p.stderr
    assert json.loads(p.stdout)["docx"].endswith(".docx")
