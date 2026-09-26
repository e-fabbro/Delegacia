"""`render`: reidentifica um produto de 04_produtos/ (fora do alcance do modelo) e grava em 04_produtos/render/.

- .md  -> .md reidentificado + .docx (títulos, parágrafos com negrito/itálico, listas, tabelas, citações,
          cabeçalho SIGILOSO e rodapé com codinome/data);
- .xlsx (matrizes) -> .xlsx reidentificado célula a célula;
- .html/.json/.txt -> mesmo formato, texto reidentificado.
Ponteiros [F:DOC-###:loc] podem ficar como estão (`manter`), virar legíveis (`legivel`, padrão) ou sair (`remover`).
"""
import datetime as dt
import re
from pathlib import Path

import docx
import openpyxl
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from agencia import caso, cofre

RE_TITULO = re.compile(r"^(#{1,6})\s+(.*)$")
RE_LISTA = re.compile(r"^\s*[-*]\s+(.*)$")
RE_LISTA_NUM = re.compile(r"^\s*\d+[.)]\s+(.*)$")
RE_CITACAO = re.compile(r"^\s*>\s?(.*)$")
RE_SEPARADOR = re.compile(r"^\|?\s*:?-{2,}")
RE_INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")
RE_PONTEIRO = re.compile(r"\[F:(DOC-\d{3,}):(p\d+|l\d+(?:-\d+)?|tx#\d+|com#\d+|ccs#\d+|ev#\d+|qsa#\d+|mov#\d+|agg#[a-z0-9_]+)\]")
FORMATOS_TEXTO = (".md", ".html", ".htm", ".json", ".txt", ".csv")


def _ponteiro_legivel(m: re.Match) -> str:
    doc, loc = m.group(1), m.group(2)
    if loc.startswith("p"):
        desc = f"pág. {loc[1:]}"
    elif loc.startswith("l"):
        desc = f"linhas {loc[1:]}"
    elif loc.startswith("tx#"):
        desc = f"lançamento {loc[3:]}"
    elif loc.startswith("com#"):
        desc = f"comunicação {loc[4:]}"
    elif loc.startswith("ccs#"):
        desc = f"registro CCS {loc[4:]}"
    elif loc.startswith("ev#"):
        desc = f"evento {loc[3:]}"
    elif loc.startswith("qsa#"):
        desc = f"registro QSA {loc[4:]}"
    elif loc.startswith("mov#"):
        desc = f"movimentação cripto {loc[4:]}"
    else:
        desc = f"análise {loc[4:]}"
    return f"[{doc}, {desc}]"


def tratar_ponteiros(texto: str, modo: str) -> str:
    if modo == "manter":
        return texto
    if modo == "remover":
        return re.sub(r"\s*" + RE_PONTEIRO.pattern, "", texto)
    return RE_PONTEIRO.sub(_ponteiro_legivel, texto)


# ---------- markdown -> docx ----------

def _celulas(linha: str) -> list[str]:
    return [c.strip() for c in linha.strip().strip("|").split("|")]


def _runs(paragrafo, texto: str) -> None:
    for parte in RE_INLINE.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
            paragrafo.add_run(parte[2:-2]).bold = True
        elif parte.startswith("`") and parte.endswith("`") and len(parte) > 2:
            r = paragrafo.add_run(parte[1:-1])
            r.font.name = "Consolas"
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            paragrafo.add_run(parte[1:-1]).italic = True
        else:
            paragrafo.add_run(parte)


def _base(doc, codinome: str) -> None:
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(11)
    secao = doc.sections[0]
    cab = secao.header.paragraphs[0]
    cab.text = "SIGILOSO — DRCC/DECOR/PCDF — uso restrito ao procedimento"
    cab.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cab.runs[0].font.size = Pt(8)
    rod = secao.footer.paragraphs[0]
    rod.text = f"Caso {codinome} · Agência Nexo · gerado em {dt.datetime.now().strftime('%d/%m/%Y %H:%M')}"
    rod.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rod.runs[0].font.size = Pt(8)


def md_para_docx(texto: str, destino: Path, codinome: str = "") -> None:
    doc = docx.Document()
    if codinome:
        _base(doc, codinome)
    linhas = texto.split("\n")
    i, em_codigo = 0, False
    while i < len(linhas):
        linha = linhas[i]
        if linha.strip().startswith("```"):
            em_codigo = not em_codigo
            i += 1
            continue
        if em_codigo:
            p = doc.add_paragraph()
            r = p.add_run(linha)
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            i += 1
            continue
        if linha.strip().startswith("|"):
            bloco = []
            while i < len(linhas) and linhas[i].strip().startswith("|"):
                if not RE_SEPARADOR.match(linhas[i].strip().strip("|").strip()):
                    bloco.append(_celulas(linhas[i]))
                i += 1
            if bloco:
                ncol = max(len(l) for l in bloco)
                tabela = doc.add_table(rows=len(bloco), cols=ncol)
                tabela.style = "Table Grid"
                for r_, l in enumerate(bloco):
                    for c, valor in enumerate(l):
                        cel = tabela.cell(r_, c)
                        cel.text = ""
                        _runs(cel.paragraphs[0], valor)
                        if r_ == 0:
                            for run in cel.paragraphs[0].runs:
                                run.bold = True
                doc.add_paragraph()
            continue
        m = RE_TITULO.match(linha)
        if m:
            doc.add_heading(re.sub(r"\*\*(.+?)\*\*", r"\1", m.group(2)).strip(), level=min(len(m.group(1)), 9))
        elif RE_LISTA.match(linha):
            _runs(doc.add_paragraph(style="List Bullet"), RE_LISTA.match(linha).group(1))
        elif RE_LISTA_NUM.match(linha):
            _runs(doc.add_paragraph(style="List Number"), RE_LISTA_NUM.match(linha).group(1))
        elif RE_CITACAO.match(linha):
            _runs(doc.add_paragraph(style="Intense Quote"), RE_CITACAO.match(linha).group(1))
        elif linha.strip() == "---":
            doc.add_page_break()
        elif linha.strip():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _runs(p, linha.strip())
        i += 1
    doc.save(destino)


# ---------- render ----------

def _xlsx(cf: cofre.Cofre, origem: Path, destino: Path, ponteiros: str) -> tuple[int, int]:
    wb = openpyxl.load_workbook(origem)
    subst = rest = 0
    for ws in wb.worksheets:
        for linha in ws.iter_rows():
            for cel in linha:
                if isinstance(cel.value, str) and cel.value:
                    novo, s, r = cf.reidentificar_contando(cel.value)
                    cel.value = tratar_ponteiros(novo, ponteiros)
                    subst += s
                    rest += r
    wb.save(destino)
    return subst, rest


def render(codinome: str, arquivo: str, ponteiros: str = "legivel", docx_: bool = True) -> dict:
    d = caso.caminho(codinome)
    base = (d / "04_produtos").resolve()
    alvo = (base / arquivo).resolve()
    if base not in alvo.parents or not alvo.is_file():
        raise caso.ErroCaso(f"arquivo {arquivo!r} não encontrado em 04_produtos/")
    if base / "render" in alvo.parents:
        raise caso.ErroCaso("o arquivo já está em 04_produtos/render/")
    if ponteiros not in ("legivel", "manter", "remover"):
        raise caso.ErroCaso("--ponteiros aceita legivel, manter ou remover")
    cf = cofre.Cofre(d / "_cofre" / "identidades.db")
    pasta = base / "render"
    pasta.mkdir(exist_ok=True)
    saida = {"codinome": codinome, "origem": str(alvo.relative_to(base.parent)), "ponteiros": ponteiros}

    if alvo.suffix.lower() == ".xlsx":
        destino = pasta / alvo.name
        subst, rest = _xlsx(cf, alvo, destino, ponteiros)
        saida["xlsx"] = str(destino)
    elif alvo.suffix.lower() in FORMATOS_TEXTO:
        texto, subst, rest = cf.reidentificar_contando(alvo.read_text(encoding="utf-8"))
        texto = tratar_ponteiros(texto, ponteiros)
        destino = pasta / alvo.name
        destino.write_text(texto, encoding="utf-8")
        saida[alvo.suffix.lstrip(".").lower()] = str(destino)
        if alvo.suffix.lower() == ".md" and docx_:
            docx_out = pasta / (alvo.stem + ".docx")
            md_para_docx(texto, docx_out, codinome)
            saida["docx"] = str(docx_out)
    else:
        raise caso.ErroCaso(f"formato não suportado pelo render: {alvo.suffix}")
    saida.update({"tokens_substituidos": subst, "tokens_remanescentes": rest, "ok": rest == 0})
    return saida
