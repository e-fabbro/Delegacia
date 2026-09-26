"""F5 — `matrizes`: consolida em 04_produtos/matrizes.xlsx (pseudonimizado) as tabelas do caso e das análises.

Uma planilha por bloco; blocos sem dados são omitidos e listados em `omitidas`. A reidentificação
(`render <COD> matrizes.xlsx`) acontece fora do modelo, em 04_produtos/render/.
"""
import json

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from agencia import achados, banco, caso, db, grafo, integracao, rif, util

CAB_FILL = PatternFill("solid", fgColor="DDDDDD")
LARGURA_MAX = 60


def _celula(v):
    if isinstance(v, (dict, list, set, tuple)):
        return json.dumps(sorted(v) if isinstance(v, set) else v, ensure_ascii=False, default=str)
    return v


def _planilha(wb, nome: str, linhas: list[dict], colunas: list[str] | None = None) -> int:
    if not linhas:
        return 0
    colunas = colunas or list(dict.fromkeys(k for l in linhas for k in l))
    ws = wb.create_sheet(nome[:31])
    ws.append(colunas)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = CAB_FILL
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for l in linhas:
        ws.append([_celula(l.get(c)) for c in colunas])
    for i, col in enumerate(colunas, start=1):
        largura = max(len(str(col)), *(len(str(_celula(l.get(col)) or "")) for l in linhas))
        ws.column_dimensions[get_column_letter(i)].width = min(LARGURA_MAX, max(8, largura + 2))
    ws.freeze_panes = "A2"
    return len(linhas)


def _tentar(fn, *a, **k):
    try:
        return fn(*a, **k)
    except caso.ErroCaso:
        return None


def _manifesto(d) -> list[dict]:
    arq = d / "01_custodia" / "manifesto.json"
    if not arq.is_file():
        return []
    return [{"doc_id": e["doc_id"], "nome": e["nome"], "tipo": e["tipo"], "sha256": e["sha256"], "tamanho": e["tamanho"],
             "recebido_em": e["recebido_em"], "paginas": e["extracao"].get("paginas"), "linhas": e["extracao"].get("linhas"),
             "tabelas": len(e["extracao"].get("tabelas", [])), "avisos": "; ".join(e.get("avisos", []))} for e in json.loads(arq.read_text(encoding="utf-8"))]


def _achados_planos(d) -> list[dict]:
    saida = []
    for agente, a in grafo._achados(d):
        saida.append({
            "agente": agente, "id": a.get("id"), "rotulo": a.get("rotulo"), "relevancia": a.get("relevancia"), "enunciado": a.get("enunciado"),
            "raciocinio": a.get("raciocinio"), "entidades": ", ".join(a.get("entidades") or []),
            "valores": "; ".join(f"{v.get('valor')} ({v.get('origem')})" for v in a.get("valores") or [] if isinstance(v, dict)),
            "periodo_inicio": (a.get("periodo") or {}).get("inicio"), "periodo_fim": (a.get("periodo") or {}).get("fim"),
            "fontes": " ".join(util.ponteiro(f["doc_id"], f["localizador"]) for f in a.get("fontes") or [] if isinstance(f, dict) and "doc_id" in f and "localizador" in f),
            "diligencia": json.dumps(a["diligencia"], ensure_ascii=False) if a.get("diligencia") else None,
        })
    return saida


def matrizes(codinome: str) -> dict:
    d = caso.caminho(codinome)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    contagens, omitidas = {}, []

    def bloco(nome, linhas, colunas=None):
        n = _planilha(wb, nome, linhas or [], colunas)
        (contagens.__setitem__(nome, n) if n else omitidas.append(nome))

    bloco("Documentos", _manifesto(d))

    r = _tentar(banco.resumo, codinome)
    bloco("Contas", [{k: v for k, v in c.items() if k not in ("maior_credito", "maior_debito")}
                     | {"maior_credito": (c.get("maior_credito") or {}).get("valor"), "maior_credito_ptr": (c.get("maior_credito") or {}).get("ponteiro"),
                        "maior_debito": (c.get("maior_debito") or {}).get("valor"), "maior_debito_ptr": (c.get("maior_debito") or {}).get("ponteiro")}
                     for c in r["contas"]] if r else [])
    lanc = _tentar(banco.lancamentos, codinome, limite=None)
    bloco("Lancamentos", lanc["lancamentos"] if lanc else [])
    r = _tentar(banco.integridade, codinome)
    if r:
        bloco("Integridade", [{"conta": c["conta"], "titular": c["titular"], "periodo_inicio": c["periodo_inicio"], "periodo_fim": c["periodo_fim"],
                              "lancamentos": c["lancamentos"], "lacunas": len(c["lacunas"]), "lacunas_detalhe": c["lacunas"],
                              "saldo_conferido": c["saldo"]["conferido"], "divergencias_saldo": c["saldo"]["divergencias"],
                              "sem_contraparte_pct": c["sem_contraparte"]["pct"], "sem_contraparte_anomalo": c["sem_contraparte"]["n_anomalo"],
                              "duplicidades": c["duplicidades"]} for c in r["contas"]]
              + [{"conta": c["conta"], "titular": c["titular"], "lancamentos": 0, "lacunas_detalhe": "CCS sem extrato"} for c in r["contas_ccs_sem_extrato"]])
    r = _tentar(banco.contrapartes, codinome, top=50)
    bloco("Contrapartes", [{"conta": c["conta"], "titular": c["titular"], **x} for c in r["contas"] for x in c["contrapartes"]] if r else [])
    r = _tentar(banco.especie, codinome)
    bloco("Especie", [{"titular": c["titular"], **l} for c in r["contas"] for l in c["lancamentos"]] if r else [])
    r = _tentar(banco.fracionamento, codinome)
    bloco("Fracionamento", r["grupos"] if r else [])
    r = _tentar(banco.passagem, codinome)
    bloco("Passagem", [{k: v for k, v in c.items() if k != "pares"} for c in r["contas"]] if r else [])
    r = _tentar(banco.circularidade, codinome)
    bloco("Circularidade", [{"ciclo": " → ".join(c["nos"] + [c["nos"][0]]), "comprimento": c["comprimento"], "valor_minimo": c["valor_minimo_no_ciclo"],
                            "ordem_temporal_coerente": c["ordem_temporal_coerente"], **p} for c in r["lista"] for p in c["passos"]] if r else [])
    r = _tentar(banco.cruzar_alvos, codinome)
    bloco("Cruzamentos", r["pares"] if r else [])
    r = _tentar(banco.linha_tempo, codinome)
    bloco("LinhaTempo_Banco", r["serie"] if r else [])

    r = _tentar(rif.comunicacoes, codinome)
    bloco("RIF_Comunicacoes", r["comunicacoes"] if r else [])
    r = _tentar(rif.envolvidos, codinome)
    bloco("RIF_Envolvidos", [{k: v for k, v in e.items() if k != "como_titular"} | {f"titular_{k}": v for k, v in e["como_titular"].items()} for e in r["lista"]] if r else [])
    r = _tentar(rif.sobreposicao, codinome)
    bloco("RIF_Sobreposicao", (r["grupos"] if r else []))
    bloco("RIF_Consolidado", (r["por_titular"] if r else []))

    con = db.abrir_caso(d)
    vinc = [dict(l) for l in con.execute("select * from vinculos order by tipo, origem, destino")]
    con.close()
    for v in vinc:
        v["total"] = util.reais(v.pop("total_centavos"))
    bloco("Vinculos", vinc)
    r = _tentar(grafo.centrais, codinome, top=500)
    bloco("Centrais", r["centrais"] if r else [])
    r = _tentar(integracao.integrada, codinome)
    bloco("LinhaTempo_Integrada", r["serie"] if r else [])
    bloco("Marcos", r["marcos"] if r else [])

    bloco("Achados", _achados_planos(d))
    r = _tentar(achados.diligencias, codinome)
    bloco("Diligencias", [{"tipo": g["tipo"], **x} for g in r["por_tipo"] for x in g["diligencias"]] if r else [])

    if not wb.sheetnames:
        raise caso.ErroCaso("nada para consolidar: ingira e importe fontes antes de `matrizes`")
    destino = d / "04_produtos" / "matrizes.xlsx"
    wb.save(destino)
    return {"codinome": codinome, "arquivo": "04_produtos/matrizes.xlsx", "planilhas": contagens, "omitidas": omitidas}
